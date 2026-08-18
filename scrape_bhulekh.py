import os
import re
import time
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import requests

INPUT_FILE = "data.xlsx"
OUTPUT_FILE = "bhulekh_scraped_output.xlsx"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    )
}


def get_text_by_id(soup, element_id):
    elem = soup.find(id=element_id)
    return elem.text.strip() if elem else ""


def clean_plot_number(raw_text):
    if not raw_text:
        return ""
    match = re.match(r"^\s*([\d/\-]+)", str(raw_text))
    return match.group(1) if match else raw_text.strip()


def parse_khata_html(html_content, url=""):
    soup = BeautifulSoup(html_content, "html.parser")
    page_text = soup.get_text()

    if "ପରିଶିଷ୍ଟ - କ" in page_text:
        khata_type = "Consolidation"
    elif "Schedule I Form No.39-A" in page_text:
        khata_type = "Settlement"
    else:
        khata_type = "Unknown"

    mouja = get_text_by_id(soup, "gvfront_ctl02_lblMouja")
    tehsil = get_text_by_id(soup, "gvfront_ctl02_lblTehsil")
    thana = get_text_by_id(soup, "gvfront_ctl02_lblThana")
    district = get_text_by_id(soup, "gvfront_ctl02_lblDist")
    khatiyan_no = get_text_by_id(soup, "gvfront_ctl02_lblKhatiyanslNo")
    tenant_name = get_text_by_id(soup, "gvfront_ctl02_lblName")
    landlord_name = get_text_by_id(soup, "gvfront_ctl02_lblLandlordName")
    status = get_text_by_id(soup, "gvfront_ctl02_lblStatua")
    publish_date = get_text_by_id(soup, "gvfront_ctl02_lblLastPublishDate")

    plot_list = []
    total_khata_acres = 0.0

    ror_table = soup.find("table", id="gvRorBack")
    if ror_table:
        rows = ror_table.find_all("tr")
        for row in rows:
            col_8_elem = row.find(id=re.compile(r"gvRorBack_ctl\d+_lblPlotNo"))
            plot_ci_elem = row.find(
                id=re.compile(r"gvRorBack_ctl\d+_lblPlotci")
            )
            plot_cni_elem = row.find(
                id=re.compile(r"gvRorBack_ctl\d+_lblPlotcni")
            )

            if col_8_elem or plot_ci_elem or plot_cni_elem:
                chaka_no = ""
                actual_plot_no = ""
                ltype = ""

                col_8_text = col_8_elem.text.strip() if col_8_elem else ""
                plot_ci_text = (
                    plot_ci_elem.text.strip() if plot_ci_elem else ""
                )
                plot_cni_text = (
                    plot_cni_elem.text.strip() if plot_cni_elem else ""
                )

                if khata_type == "Consolidation":
                    if plot_ci_text:
                        actual_plot_no = clean_plot_number(plot_ci_text)
                        chaka_no = clean_plot_number(col_8_text)
                        ltype_elem = row.find(
                            id=re.compile(r"gvRorBack_ctl\d+_lbllType")
                        )
                        ltype = ltype_elem.text.strip() if ltype_elem else ""
                    elif plot_cni_text:
                        actual_plot_no = clean_plot_number(plot_cni_text)
                        chaka_no = ""
                        cni_elem = row.find(
                            id=re.compile(r"gvRorBack_ctl\d+_lblCNItype")
                        )
                        ltype = cni_elem.text.strip() if cni_elem else ""
                else:
                    actual_plot_no = clean_plot_number(col_8_text)
                    chaka_no = ""
                    ltype_elem = row.find(
                        id=re.compile(r"gvRorBack_ctl\d+_lbllType")
                    )
                    ltype = ltype_elem.text.strip() if ltype_elem else ""

                if actual_plot_no:
                    acre_elem = row.find(
                        id=re.compile(r"gvRorBack_ctl\d+_lblAcre")
                    )
                    dec_elem = row.find(
                        id=re.compile(r"gvRorBack_ctl\d+_lblDecimil")
                    )

                    acre_str = acre_elem.text.strip() if acre_elem else "0"
                    dec_str = dec_elem.text.strip() if dec_elem else "0"

                    try:
                        x_acre = float(acre_str) if acre_str else 0.0
                    except ValueError:
                        x_acre = 0.0

                    try:
                        yyyy_dec = float(dec_str) if dec_str else 0.0
                    except ValueError:
                        yyyy_dec = 0.0

                    calculated_acres = round(x_acre + (yyyy_dec / 10000.0), 4)
                    total_khata_acres += calculated_acres

                    entry_parts = [f"Plot: {actual_plot_no}"]
                    if chaka_no:
                        entry_parts.append(f"Chaka: {chaka_no}")
                    if ltype:
                        entry_parts.append(f"Kisam: {ltype}")
                    entry_parts.append(f"Area: {calculated_acres} Acres")

                    plot_list.append(f"[{' | '.join(entry_parts)} ]")

    total_plots = len(plot_list)
    plots_formatted = (
        ",\n".join(plot_list) if total_plots > 0 else "No Plots Available"
    )

    return {
        "URL": url,
        "Khata_Type": khata_type,
        "Khatiyan_No": khatiyan_no,
        "Mouja": mouja,
        "Tehsil": tehsil,
        "Thana": thana,
        "District": district,
        "Tenant_Name": tenant_name,
        "Landlord_Name": landlord_name,
        "Status": status,
        "Publish_Date": publish_date,
        "Total_Plots_Count": total_plots,
        "Total_Khata_Area_Acres": round(total_khata_acres, 4),
        "Plot_Details": plots_formatted,
    }


def main():
    print("=" * 65)
    print("Bhulekh Odisha Auto-URL Generator & RoR Scraper")
    print("=" * 65)

    if not os.path.exists(INPUT_FILE):
        print(f"[ERROR] '{INPUT_FILE}' not found in current folder.")
        return

    df_input = pd.read_excel(INPUT_FILE)
    print(f"Loaded '{INPUT_FILE}' with {len(df_input)} rows.\n")

    col_dist_code = df_input.columns[1]
    col_teh_code = df_input.columns[3]
    col_vill_code = df_input.columns[5]
    col_khata_no = (
        df_input.columns[6] if len(df_input.columns) > 6 else "KhataNo"
    )

    session = requests.Session()
    session.headers.update(HEADERS)
    results = []

    for idx, row in df_input.iterrows():
        try:
            dist_code = int(row[col_dist_code])
            teh_code = int(row[col_teh_code])
            vill_code = int(row[col_vill_code])
            khata_val = row.get(col_khata_no, 1)
            khata_no = int(khata_val) if pd.notna(khata_val) else 1

            target_url = (
                f"http://bhulekh.ori.nic.in/ViewRoR.aspx?"
                f"DistCode={dist_code}&TehCode={teh_code}&VillCode={vill_code}&KhataNo={khata_no}&type=front"
            )
        except Exception as e:
            print(f"[{idx+1}/{len(df_input)}] Error reading row codes: {e}")
            continue

        print(f"[{idx+1}/{len(df_input)}] Scraping: {target_url}")

        try:
            response = session.get(target_url, timeout=20)
            if response.status_code == 200:
                parsed = parse_khata_html(response.text, url=target_url)
                results.append(parsed)
                print(
                    f"   -> Success: [{parsed['Khata_Type']}] "
                    f"Khata: {parsed['Khatiyan_No']}, Plots: {parsed['Total_Plots_Count']}"
                )
            else:
                results.append(
                    {"URL": target_url, "Khata_Type": f"HTTP {response.status_code}"}
                )
        except Exception as e:
            print(f"   -> Timeout / Network Error: {e}")
            results.append({"URL": target_url, "Khata_Type": "Error / Timeout"})

        time.sleep(0.1)

    print(f"\nWriting output to '{OUTPUT_FILE}'...")
    output_df = pd.DataFrame(results)

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        output_df.to_excel(writer, index=False, sheet_name="Scraped_Data")
        wb = writer.book
        ws = writer.sheets["Scraped_Data"]

        header_fill = PatternFill(
            start_color="1F497D", end_color="1F497D", fill_type="solid"
        )
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )

        for row in ws.iter_rows(
            min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column
        ):
            for cell in row:
                cell.border = thin_border
                if cell.column == 14:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(vertical="top")

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            if col_letter == "N":
                ws.column_dimensions[col_letter].width = 65
            else:
                ws.column_dimensions[col_letter].width = min(
                    max(max_len + 3, 12), 35
                )

    print("=" * 65)
    print(f"Scraping completed! Results saved to '{OUTPUT_FILE}'.")
    print("=" * 65)


if __name__ == "__main__":
    main()
